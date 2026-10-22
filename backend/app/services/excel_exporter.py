"""Professional multi-tab .xlsx / .csv exporter (openpyxl, Decimal-native)."""

from __future__ import annotations

import csv
from decimal import Decimal
from io import BytesIO, StringIO
from typing import TypeAlias

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.security import protect_cell

CellValue: TypeAlias = str | int | Decimal
Row: TypeAlias = dict[str, CellValue]

AUDIT_NOTE = (
    "Generado por VolleyOS Global (AGPL-3.0-or-later). Cálculos en Decimal, "
    "procesados localmente, sin telemetría. Este reporte es informativo y no "
    "constituye asesoramiento médico ni decisión arbitral oficial."
)
HEADER_FILL = PatternFill("solid", fgColor="E11D48")
TITLE_FILL = PatternFill("solid", fgColor="070709")
NOTE_FILL = PatternFill("solid", fgColor="FEF3C7")


def _width(value: CellValue) -> int:
    return min(50, max(12, len(str(value)) + 2))


def export_xlsx(
    kind: str,
    title: str,
    summary: dict[str, CellValue],
    rows: list[Row],
    audit_note: str = AUDIT_NOTE,
) -> bytes:
    if not title.strip():
        raise ValueError("title is required")
    headers = list(rows[0].keys()) if rows else []
    if rows and any(list(row.keys()) != headers for row in rows):
        raise ValueError("all detail rows must share the same columns")

    workbook = Workbook()
    workbook.properties.creator = "VolleyOS Global"
    workbook.properties.title = title
    workbook.properties.description = f"VolleyOS {kind} report"
    workbook.properties.license = "AGPL-3.0-or-later"

    summary_sheet = workbook.active
    assert summary_sheet is not None
    summary_sheet.title = "Resumen"
    summary_sheet["A1"] = title
    summary_sheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    summary_sheet["A1"].fill = TITLE_FILL
    summary_sheet.merge_cells("A1:D1")
    summary_sheet["A2"] = f"Tipo: {kind}"
    summary_sheet["A2"].font = Font(italic=True, color="6B7280")

    summary_sheet["A4"] = "Métrica"
    summary_sheet["B4"] = "Valor"
    for cell in ("A4", "B4"):
        summary_sheet[cell].font = Font(bold=True, color="FFFFFF")
        summary_sheet[cell].fill = HEADER_FILL
    row_index = 5
    for key, value in summary.items():
        summary_sheet.cell(row=row_index, column=1, value=key)
        target = summary_sheet.cell(row=row_index, column=2, value=protect_cell(value))
        if isinstance(value, Decimal):
            target.number_format = "0.00"
        elif isinstance(value, int):
            target.number_format = "#,##0"
        target.alignment = Alignment(horizontal="left")
        row_index += 1
    summary_sheet.column_dimensions["A"].width = 34
    summary_sheet.column_dimensions["B"].width = 28

    note_row = row_index + 1
    summary_sheet.cell(row=note_row, column=1, value=audit_note)
    summary_sheet.cell(row=note_row, column=1).fill = NOTE_FILL
    summary_sheet.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    summary_sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=4)
    summary_sheet.row_dimensions[note_row].height = 48

    detail = workbook.create_sheet("Detalle")
    for column, header in enumerate(headers, start=1):
        cell = detail.cell(row=1, column=column, value=protect_cell(header))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for row_number, row in enumerate(rows, start=2):
        for column, header in enumerate(headers, start=1):
            value = row[header]
            cell = detail.cell(row=row_number, column=column, value=protect_cell(value))
            if isinstance(value, Decimal):
                cell.number_format = "0.00"
            elif isinstance(value, int):
                cell.number_format = "#,##0"
    detail.freeze_panes = "A2"
    for column, header in enumerate(headers, start=1):
        widest = max([_width(header)] + [_width(row[header]) for row in rows])
        detail.column_dimensions[get_column_letter(column)].width = widest

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_csv(rows: list[Row]) -> bytes:
    if not rows:
        raise ValueError("at least one row is required")
    headers = list(rows[0].keys())
    buffer = StringIO()
    safe_headers = [str(protect_cell(header)) for header in headers]
    writer = csv.DictWriter(buffer, fieldnames=safe_headers, lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow(
            {
                safe: protect_cell(row[header])
                for safe, header in zip(safe_headers, headers, strict=True)
            }
        )
    return buffer.getvalue().encode("utf-8")
