from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.excel_exporter import AUDIT_NOTE, export_csv, export_xlsx


def test_xlsx_has_summary_and_detail() -> None:
    payload = export_xlsx(
        kind="scouting",
        title="Scouting — Final",
        summary={"Ataques": 24, "Eficiencia": Decimal("0.42")},
        rows=[
            {"Jugador": 7, "Zona": 4, "Resultado": "KILL", "Eficiencia": Decimal("0.50")},
            {"Jugador": 9, "Zona": 2, "Resultado": "ERROR", "Eficiencia": Decimal("-0.10")},
        ],
    )
    workbook = load_workbook(BytesIO(payload))
    assert workbook.sheetnames == ["Resumen", "Detalle"]
    summary = workbook["Resumen"]
    assert summary["A1"].value == "Scouting — Final"
    assert summary["A5"].value == "Ataques"
    assert summary["B6"].value == 0.42
    detail = workbook["Detalle"]
    assert detail["A1"].value == "Jugador"
    assert detail["A3"].value == 9
    assert detail.freeze_panes == "A2"
    notes = [cell.value for row in summary.iter_rows() for cell in row if cell.value]
    assert AUDIT_NOTE in notes


def test_xlsx_without_rows() -> None:
    payload = export_xlsx(kind="generic", title="Vacío", summary={"Total": 0}, rows=[])
    workbook = load_workbook(BytesIO(payload))
    assert workbook["Detalle"].max_row == 1


def test_xlsx_rejects_ragged_rows() -> None:
    with pytest.raises(ValueError, match="same columns"):
        export_xlsx(
            kind="scouting",
            title="Mal",
            summary={},
            rows=[{"A": 1}, {"B": 2}],
        )


def test_xlsx_requires_title() -> None:
    with pytest.raises(ValueError, match="title"):
        export_xlsx(kind="generic", title="  ", summary={}, rows=[])


def test_csv_export() -> None:
    payload = export_csv([{"Jugador": 7, "Eficiencia": Decimal("0.50")}])
    assert payload.decode("utf-8").startswith("Jugador,Eficiencia")
    assert "0.50" in payload.decode("utf-8")


def test_csv_empty_rejected() -> None:
    with pytest.raises(ValueError, match="at least one row"):
        export_csv([])
